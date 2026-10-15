"""docs/セラゾイド変調部の動作説明.docx を作る（python-docx）。図は make_serrasoid_figs.py で先に作っておく"""
import os, datetime
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(DOCS, 'images')
OUT = os.path.join(DOCS, 'セラゾイド変調部の動作説明.docx')
JP = '游ゴシック'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = sec.bottom_margin = Mm(18)

def set_font(style, size, bold=None, color=None):
    style.font.size = Pt(size); style.font.name = JP
    rpr = style.element.get_or_add_rPr(); f = rpr.find(qn('w:rFonts'))
    if f is None: f = OxmlElement('w:rFonts'); rpr.append(f)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): f.set(qn(k), JP)
    if bold is not None: style.font.bold = bold
    if color is not None: style.font.color.rgb = color
set_font(doc.styles['Normal'], 10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Normal'].paragraph_format.line_spacing = 1.25
for name, size in (('Title', 20), ('Heading 1', 14), ('Heading 2', 12)):
    set_font(doc.styles[name], size, True, RGBColor(0x1F, 0x3A, 0x5F))
for name in ('List Bullet', 'List Number'):
    set_font(doc.styles[name], 10.5)

def p(text, bold=False, size=None, align=None, style=None):
    para = doc.add_paragraph(style=style)
    r = para.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    if align: para.alignment = align
    return para
def bullets(items):
    for t in items: doc.add_paragraph(t, style='List Bullet')
def numbered(items):
    for t in items: doc.add_paragraph(t, style='List Number')
def formula(text):
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = para.add_run(text); r.font.size = Pt(11); r.italic = True
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)
def table(header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); shade(c, 'DCE6F1')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Mm(w)
    doc.add_paragraph()
    return t
def figure(name, caption, width=170):
    doc.add_picture(os.path.join(IMG, name), width=Mm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = p(caption, size=9, align=WD_ALIGN_PARAGRAPH.CENTER); c.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# ---------------------------------------------------------------- 表紙の見出し
doc.add_paragraph('Serrasoid TRX 29M　セラゾイド変調部の動作説明', style='Title')
p('29MHz FM 10W トランシーバ（リポジトリ Serrasoid_TRX_29M）　%s 版　JA6FUF' % datetime.date.today().strftime('%Y-%m-%d'), size=9.5)
p('対象: 回路図のブロック「セラゾイド変調」「逓倍」「送信音声」。部品番号・定数は回路図 Ver 1.0（tools/circuit.py）に合わせている。'
  '実機は未製作で、数値は計算値。', size=9.5)

# ---------------------------------------------------------------- 1
doc.add_heading('1. 概要', level=1)
p('セラゾイド変調部は、915.625kHz ののこぎり波としきい値（音声）を比べ、方形波のエッジの時刻を音声で前後させることで、'
  '位相変調（PM）をかけた信号を作る。これを 32 倍に逓倍して 29.3MHz の送信信号にする。')
bullets(['のこぎり波の傾きが一定なので、しきい値の変化がそのままエッジの時刻のずれになり、位相が音声電圧に比例する（直線性のよい PM）',
         '1 段で得られる位相のずれは小さい（最大 ±0.1rad 程度）。逓倍すると位相も 32 倍になるので、29.3MHz で FM 無線機として十分な偏移（±5kHz）になる',
         '回路は「基準クロック → のこぎり波 → 比較 → パルス整形 → 逓倍」の順に一方向へ流れる。のこぎり波と比較までが変調、そのあとは波形を整えて周波数を上げる部分',
         '1940 年代後半の放送用 FM 送信機の方式を、汎用のロジック IC・コンパレータ・トランジスタで組み直したもの'])
figure('serrasoid_flow.png', '図 1　セラゾイド変調部の信号の流れ（赤字はネット名、TP は観測点）')

# ---------------------------------------------------------------- 2
doc.add_heading('2. 基準クロックとリセットパルス', level=1)
p('Si5351 の CLK0（9.15625MHz）を 74HCT4017 で 10 分周し、1.092µs（915.625kHz）ごとに幅 109ns のリセットパルスを作る。'
  'このパルスが、のこぎり波を 1 周期ごとに 0V に戻す。')
table(['部品', 'はたらき'], [
    ['U6 Si5351（秋月 AE-Si5351A）', 'CLK0 = 9.15625MHz。送信中だけ出す。チャンネルは CLK0 を 3.125kHz ずつ変える（出力で 10kHz ステップ）'],
    ['R11 100Ω', 'CLK0 の直列抵抗。反射・リンギングを抑える'],
    ['U7 74HCT4017', 'ジョンソンカウンタ（÷10）。Q0 が 10 クロックに 1 回、1 クロック分（109ns）だけ High。デコード出力にグリッチが出ない。'
                     'HCT 品なので Si5351 の 3.3V 出力をそのまま受けられる（High 入力 2.0V 以上）'],
    ['R14 100Ω、Q3 2N7000', 'Q0 が High の間だけオンし、のこぎり波のコンデンサ C19 を放電する。オン抵抗は数 Ω なので 109ns の間に 0V まで下がる'],
], [55, 115])
p('受信中は CLK0 を止めるので変調部は動かない（自局の信号が受信部に回り込まない）。')

# ---------------------------------------------------------------- 3
doc.add_heading('3. のこぎり波の生成', level=1)
p('2SA1015 の定電流源で C19（100pF）を一定の電流で充電し、直線的に上がるのこぎり波 RAMP を作る。傾きは約 2.5V/µs、山の高さは約 2.4V。')
table(['部品', 'はたらき'], [
    ['D4・D5 1N4148、R13 10kΩ', 'ベースの基準電圧。+5VT からダイオード 2 本分（約 1.3V）下の 3.7V に固定する'],
    ['C18 0.1µF', '+5VT とベースの間。電源の変動にベースが追従し、電流が一定に保たれる'],
    ['R12 1.5kΩ ＋ VR1 1kΩ', 'エミッタ抵抗。両端の電圧（約 0.65V）÷ 抵抗 で電流が決まる。VR1 で傾きを調整'],
    ['Q2 2SA1015', 'PNP の定電流源。コレクタから C19 へ電流を流す'],
    ['C19 100pF（C0G）', 'のこぎり波のコンデンサ。温度で容量が変わると傾きが変わるので C0G（NP0）を使う'],
    ['TP1', 'のこぎり波の観測点'],
], [55, 115])
formula('充電電流  I ＝ 0.65V ÷（R12 ＋ VR1）＝ 0.65 ÷ 1.5k〜2.5k ＝ 260〜430µA')
formula('傾き  dV/dt ＝ I ÷ C ≒ 300µA ÷ 120pF ≒ 2.5V/µs')
formula('山の高さ  ≒ 2.5V/µs ×（1.092µs − 0.109µs）≒ 2.4V')
p('C は C19 の 100pF に、2N7000 の出力容量・MAX913 の入力容量・配線の容量（合わせて約 20pF）を足した値。'
  '定電流源が飽和しない上限はコレクタ電圧で約 4.0V なので、2.4V までは傾きが一定に保たれる。')

# ---------------------------------------------------------------- 4
doc.add_heading('4. しきい値と音声（送信音声ブロック）', level=1)
p('しきい値 THRESH は「直流の中心（約 1.2V）＋ 音声」で、MCP6004 の 4 回路で作る。音声が大きくなっても ±0.6V でクリップし、'
  '3kHz の LPF を通してから変調部へ入れるので、偏移と占有帯域が一定以内に収まる。')
table(['段', '部品', '内容'], [
    ['マイク', 'J7、R24・C33・R25', 'エレクトレットマイクに 5V（+5VT）から約 2.2kΩ でバイアス。約 10mVp の音声'],
    ['A マイクアンプ', 'U13A、R26・R27・C35', '反転増幅 ×22（220k / 10k）。220pF で高域を落とす'],
    ['マイク利得', 'VR3', '次段に入れる大きさを調整'],
    ['B 利得段', 'U13B、R28・R29', '非反転増幅 ×11（1 + 100k / 10k）'],
    ['IDC（クリップ）', 'R30・D6・D7', 'VREF（2.5V）を中心に ±0.6V で頭を切る。大声でも偏移が一定以上にならない'],
    ['C 3kHz LPF', 'U13C、R31・R32・C37・C38', 'サレンキー 2 次。fc ≒ 3.3kHz、Q ≒ 0.73。クリップでできた高調波を落とす'],
    ['偏移', 'VR4、R33 100kΩ', '音声の大きさを決める（±5kHz に合わせる）'],
    ['中心', 'VR5、C40、R34 10kΩ', 'しきい値の直流の中心。エッジがのこぎり波の中央に来るように合わせる'],
    ['D バッファ', 'U13D', 'ボルテージフォロワ。THRESH を低インピーダンスで MAX913 に渡す'],
], [28, 45, 97])
formula('THRESH ＝（VR5 の電圧 × 100k ＋ VR4 の音声 × 10k）÷ 110k　→　音声は最大 ±0.6V × 10/110 ≒ ±55mV')
p('C40（10µF）で VR5 側は交流的に GND になるので、音声は R33 と R34 の比（10k : 110k）で約 1/11 に下がる。2.5V の基準電圧 VREF は'
  '10kΩ ＋ 10kΩ の分圧と 10µF（C32）で作る。')

# ---------------------------------------------------------------- 5
doc.add_heading('5. コンパレータとパルス位置変調', level=1)
p('MAX913（U8）が RAMP（＋入力）と THRESH（−入力）を比べ、RAMP が THRESH を超えた瞬間に出力 COMP_OUT を High にする。'
  'しきい値が上がると交点が遅れ、下がると早まる。この時刻のずれがそのまま位相のずれになる。')
figure('serrasoid_waves.png', '図 2　のこぎり波としきい値、コンパレータとワンショットの出力（音声の振れは見やすく誇張）')
formula('時刻のずれ  Δt ＝ ΔV ÷（dV/dt）')
formula('位相のずれ  Δφ ＝ 2π × Δt ÷ T　（T ＝ 1.092µs）')
p('たとえば ΔV ＝ 1mV なら Δt ＝ 0.4ns、Δφ ＝ 2.3mrad。のこぎり波の傾きが一定なので、Δφ は ΔV に比例する。')
bullets(['MAX913 は単電源 5V（V− を GND）で使う。入力範囲は −0.2〜+3.5V で、のこぎり波（0〜2.4V）としきい値（約 1.2V）が入る',
         'LE（5 番ピン）は GND。Low でラッチせず比較を続ける（High または開放でラッチ）',
         '出力は TTL レベル（High 2.4〜3.4V）なので、次段は TTL レベルを受けられる 74HCT123 にしている',
         'コンパレータの出力は、交点で立ち上がり、次のリセット（のこぎり波が 0V に戻る）で立ち下がる。幅は音声で変わるので、次の段で幅をそろえる'])

# ---------------------------------------------------------------- 6
doc.add_heading('6. パルス整形（ワンショット）', level=1)
p('74HCT123（U9）が COMP_OUT の立ち上がりで、幅 546ns（1 周期の半分）のパルスを出す。出力 PM_OUT はデューティ 50% の方形波で、'
  '立ち上がり・立ち下がりの両方が音声で同じだけ前後する。')
table(['部品', 'はたらき'], [
    ['U9 74HCT123（1 回路目）', 'B 入力（2 番）の立ち上がりでトリガ。A（1 番）は GND、Clr（3 番）は High。2 回路目は入力を固定して使わない'],
    ['C21 100pF、R15 4.7kΩ ＋ VR2 10kΩ', '時定数。約 300〜760ns の範囲で調整し、546ns（デューティ 50%）に合わせる'],
    ['TP3', 'PM_OUT の観測点。立ち上がりが音声で揺れて見える'],
], [55, 115])
p('デューティを 50% にするのは、次の逓倍で各段の出力をきれいな 2 倍にするため。50% からずれると、もとの周波数の成分が残る（後段の BPF で落とす）。')

# ---------------------------------------------------------------- 7
doc.add_heading('7. 逓倍 ×32', level=1)
p('XOR（74AC86）による 2 倍の段を 5 つ重ね、915.625kHz を 29.3MHz にする。各段は XOR（x, x を遅らせたもの）で、入力の両方のエッジで'
  '遅延 τ の幅のパルスが出るので、周波数が 2 倍になる。')
figure('serrasoid_doubler.png', '図 3　XOR による 2 倍の段（RC ＋ 74AC14 で τ ＝ 入力周期の 1/4 遅らせる）')
table(['段', '入力 → 出力', '遅延 τ（入力周期の 1/4）', 'RC', '部品'], [
    ['1', '0.916 → 1.831MHz', '273ns', '3.9kΩ・100pF', 'R16・C23・U10A・U11A'],
    ['2', '1.831 → 3.663MHz', '137ns', '1.8kΩ・100pF', 'R17・C24・U10B・U11B'],
    ['3', '3.663 → 7.325MHz', '68ns', '910Ω・100pF', 'R18・C25・U10C・U11C'],
    ['4', '7.325 → 14.65MHz', '34ns', '510Ω・82pF', 'R19・C26・U10D・U11D'],
    ['5', '14.65 → 29.30MHz', '17ns', '330Ω・47pF', 'R20・C27・U10E・U12A'],
], [10, 38, 38, 32, 52])
bullets(['遅延は RC ＋ 74AC14 の遅れ（約 6ns）。τ を入力周期の 1/4 にすると出力のデューティが 50% になる',
         'RC のなまった波形を 74AC86 に直接入れると、AC 品は入力の立ち上がりが遅いと発振・誤動作する。そのためシュミットの 74AC14 で整形する（反転しても XOR の出力は 2 倍のまま）',
         '74AC14 のしきい値は電源の約 0.6 倍（上がるとき）と約 0.4 倍（下がるとき）。この値と 0.69RC ＝ τ の組み合わせで、波形が安定したあとの遅延がちょうど τ になる。しきい値は個体差があるので、デューティは 50% から少しずれることがある',
         '出力は U12B（バッファ）→ R21 100Ω → C31 → 送信RF の BPF のタップ。TP4 で 29.3MHz を観測できる',
         '逓倍は時刻のずれ Δt をそのまま運ぶ。周期が 1/2 になるごとに同じ Δt が 2 倍の位相になるので、5 段で位相のずれは 32 倍になる'])
formula('29.3MHz での位相のずれ ＝ 32 × Δφ')

# ---------------------------------------------------------------- 8
doc.add_heading('8. 偏移の計算例と、PM ＝ プリエンファシス付き FM', level=1)
p('変調部で ±0.052rad（しきい値で約 ±23mV）振ると、29.3MHz で ±1.67rad になり、3kHz の音で ±5kHz の偏移になる。')
formula('周波数偏移  Δf ＝ 32 × Δφ × fm　（fm ＝ 音の周波数）')
table(['音の周波数 fm', '29.3MHz での位相', '周波数偏移 Δf'], [
    ['300Hz', '±1.67rad', '±0.5kHz'],
    ['1kHz', '±1.67rad', '±1.7kHz'],
    ['3kHz', '±1.67rad', '±5.0kHz'],
], [45, 55, 55])
bullets(['PM では、同じ大きさの音なら位相のずれは音の高さによらず一定で、周波数偏移は音の高さに比例する（6dB/oct で増える）',
         'アマチュア無線の FM 機は、送信で 6dB/oct のプリエンファシス（高い音ほど強く）をかけ、受信でディエンファシスで戻す。PM はこのプリエンファシス付き FM と同じ特性になるので、積分器を入れずに PM のまま送れば、市販の FM 機で自然な音で聞こえる',
         'セラゾイド変調の本来の使い方（放送用 FM）では、音声を先に積分（1/f）して PM を FM に直していた。この機械はその積分を入れていない',
         '最大の偏移はクリップ（±0.6V）と VR4 で決まる。3kHz の LPF より上の音は入らないので、偏移は ±5kHz を超えない'])

# ---------------------------------------------------------------- 9
doc.add_heading('9. 調整手順', level=1)
p('オシロスコープと受信機（FM が受けられるもの）で、次の順に合わせる。送信は必ずダミーロードにつないで行う。')
numbered(['VR1（傾き）: TP1 ののこぎり波の山を約 2.4V にする（2.0〜2.8V なら可。3.5V を超えないこと）',
          'VR5（中心）: 音声なしで TP2 の直流を、のこぎり波の山の約半分（約 1.2V）にする。TP3 の立ち上がりが 1 周期のほぼ中央に来る',
          'VR2（パルス幅）: TP3（PM_OUT）のデューティを 50%（High 546ns）にする',
          'TP4: 29.3MHz の方形波が出ていることを確かめる。デューティが 50% から大きくずれる段があれば、その段の R を 1 本上下の値に替える',
          'VR3（マイク利得）・VR4（偏移）: 3kHz の音を入れ、受信機の偏移表示またはスペクトラムで ±5kHz に合わせる',
          '偏移の確かめ方（ベッセルヌル）: PM では変調指数 β ＝ 29.3MHz での位相のずれで、音の高さによらない。クリップするくらいの大きさの音を入れ、VR4 を上げていってキャリアが消える点（β ＝ 2.405rad）を見つける。そこから VR4 を 0.69 倍（1.67 ÷ 2.405）に戻すと、3kHz で ±5kHz になる'])

# ---------------------------------------------------------------- 10
doc.add_heading('10. 関係する部品と注意点', level=1)
table(['ブロック', '主な部品'], [
    ['セラゾイド変調', 'U6 Si5351、U7 74HCT4017、Q2 2SA1015、Q3 2N7000、C19 100pF（C0G）、U8 MAX913、U9 74HCT123、VR1・VR2、TP1〜TP3'],
    ['逓倍', 'U10 74AC14、U11・U12 74AC86、R16〜R20・C23〜C27（遅延）、R21・C31、TP4'],
    ['送信音声', 'U13 MCP6004、VR3〜VR5、D6・D7（クリップ）、J7（マイク）'],
    ['電源', '+5VT（L7805、変調部専用のリニア電源。DC-DC のリプルがしきい値に乗ると位相のゆれになるため）'],
], [35, 135])
bullets(['74AC86・74AC14 は必ず AC 品を使う（HC・HCT では 29.3MHz の段が間に合わない）',
         'C19・遅延用のコンデンサは C0G（NP0）。温度で値が変わると傾き・遅延が変わる',
         'MAX913 にはヒステリシスがない。のこぎり波は速くしきい値を通り過ぎるので通常は問題ないが、出力がばたつく場合は小さな正帰還を足す',
         '逓倍は位相雑音も 32 倍（30dB）にする。+5VT・CLK0 の配線は短くし、パスコン（C17・C20・C22・C28〜C30）は各 IC の電源ピンのそばに付ける',
         '逓倍後には分数調波（14.65MHz など）や隣の逓倍成分が残る。送信RF の BPF・LPF で落とし、スペクトラムアナライザでスプリアスを確かめる'])

# 標準テンプレートの <w:zoom> に percent がないので付ける（規格では必須）
z = doc.settings.element.find(qn('w:zoom'))
if z is not None and z.get(qn('w:percent')) is None: z.set(qn('w:percent'), '100')
doc.save(OUT)
print('wrote', OUT)
